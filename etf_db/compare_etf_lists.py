"""
MoneyDJ 台股ETF受益人數排行 vs 本機DB ETF清單 比對工具

用法：
  python etf_db/compare_etf_lists.py

輸出：~/Desktop/比對etf.xlsx（可用環境變數 COMPARE_OUT 覆蓋路徑）

工作表：
  MoneyDJ清單  — MoneyDJ Rank0016 受益人數排行（台股ETF，儘量抓齊）
  本機DB清單   — etf_master.db 的 etf_profile（607 檔）
  比對結果     — 三類：兩邊都有 / 只在MoneyDJ / 只在DB
"""

import os
import re
import sqlite3
import time
from pathlib import Path

import requests
from bs4 import BeautifulSoup

try:
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.cell.cell import MergedCell
except ImportError:
    print("請安裝：pip install openpyxl")
    raise

# ---------------------------------------------------------------------------
# 設定
# ---------------------------------------------------------------------------
DB_PATH  = Path(os.environ.get("ETF_DB_PATH",  Path.home() / "Desktop" / "etf_master.db"))
OUT_PATH = Path(os.environ.get("COMPARE_OUT",  Path.home() / "Desktop" / "比對etf.xlsx"))

RANK_URL = "https://www.moneydj.com/etf/x/Rank/Rank0016.xdjhtm"

HEADERS = {
    "User-Agent": (
        "Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) "
        "AppleWebKit/537.36 (KHTML, like Gecko) Chrome/124.0.0.0 Safari/537.36"
    ),
    "Referer": "https://www.moneydj.com/ETF/",
    "Accept-Language": "zh-TW,zh;q=0.9",
}

# MoneyDJ 受益人數排行的 eOrd 代碼
EORD_BENEFICIARY = "T100100"   # 受益人數（降冪）

# 多種排序維度：用不同排序抓前100，合併去重以突破每次100筆上限
SORT_COMBOS = [
    # (eOrd,          eSort, 說明)
    ("T100100",       "2",   "受益人數↓"),
    ("T100100",       "1",   "受益人數↑"),
    ("T100090",       "2",   "規模↓"),
    ("T100090",       "1",   "規模↑"),
    ("T100160",       "2",   "費用率↓"),
    ("T100160",       "1",   "費用率↑"),
    ("T100010",       "2",   "近一週↓"),
    ("T100010",       "1",   "近一週↑"),
    ("T100020",       "2",   "近一月↓"),
    ("T100030",       "2",   "近三月↓"),
    ("T100040",       "2",   "近六月↓"),
    ("T100050",       "2",   "近一年↓"),
    ("T100050",       "1",   "近一年↑"),
    ("T100060",       "2",   "今年↓"),
    ("T100060",       "1",   "今年↑"),
]

# eRank 分類篩選（台股ETF各子類）
ERANK_VALUES = ["allex", "twex", "alltw", "tw", "allbond", "bond",
                "allleverage", "leverage", "reverse", "reit", "active"]


# ---------------------------------------------------------------------------
# MoneyDJ 抓取
# ---------------------------------------------------------------------------
def _parse_rank_table(soup: BeautifulSoup):
    """解析最大資料表，回傳 (欄位清單, 資料列清單)"""
    best_table = None
    best_count = 0
    for t in soup.find_all("table"):
        data_rows = [tr for tr in t.find_all("tr") if tr.find("td")]
        if len(data_rows) > best_count:
            best_table = t
            best_count = len(data_rows)
    if not best_table or best_count < 3:
        return [], []

    headers = [th.get_text(strip=True) for th in best_table.find_all("th")]
    rows = []
    for tr in best_table.find_all("tr"):
        cells = [td.get_text(strip=True) for td in tr.find_all("td")]
        if not cells:
            continue
        if headers and len(headers) == len(cells):
            rows.append(dict(zip(headers, cells)))
        else:
            rows.append({f"col{i}": v for i, v in enumerate(cells)})
    return headers, rows


def _extract_code(row: dict) -> str:
    """從一列資料找 ETF 代碼欄位"""
    for v in row.values():
        if re.match(r"^\d{4,6}[A-Z]{0,2}$", str(v).strip()):
            return str(v).strip()
    return ""


def fetch_moneydj_etfs() -> list:
    """
    用多種排序方向各抓 100 筆，合併去重，盡量湊出 MoneyDJ 所有台股ETF。
    每次請求間隔 0.35 秒（有禮貌限速）。
    """
    s = requests.Session()
    s.headers.update(HEADERS)
    all_rows_by_code: dict = {}   # code → 第一次見到的那列資料

    total_requests = 0

    # 策略 A：預設頁（不帶參數，可能一次全列）
    try:
        r = s.get(RANK_URL, timeout=15)
        r.encoding = "utf-8"
        _, rows = _parse_rank_table(BeautifulSoup(r.text, "lxml"))
        for row in rows:
            code = _extract_code(row)
            if code and code not in all_rows_by_code:
                all_rows_by_code[code] = row
        total_requests += 1
        print(f"  [預設頁] {len(rows)} 筆 → 累計 {len(all_rows_by_code)} 檔")
        time.sleep(0.35)
    except Exception as e:
        print(f"  [預設頁] 錯誤：{e}")

    # 策略 B：不同 eRank 分類
    for erank in ERANK_VALUES:
        try:
            r = s.get(RANK_URL,
                      params={"eRank": erank, "eOrd": EORD_BENEFICIARY, "eSort": "2"},
                      timeout=12)
            r.encoding = "utf-8"
            _, rows = _parse_rank_table(BeautifulSoup(r.text, "lxml"))
            new = 0
            for row in rows:
                code = _extract_code(row)
                if code and code not in all_rows_by_code:
                    all_rows_by_code[code] = row
                    new += 1
            total_requests += 1
            if new > 0 or len(rows) > 0:
                print(f"  [eRank={erank:12s}] {len(rows)} 筆, 新增 {new} → 累計 {len(all_rows_by_code)}")
            time.sleep(0.35)
        except Exception as e:
            print(f"  [eRank={erank}] 錯誤：{e}")

    # 策略 C：不同排序維度（allex 分類）
    before_c = len(all_rows_by_code)
    for eord, esort, label in SORT_COMBOS:
        if len(all_rows_by_code) >= 400:   # 夠了就停
            break
        try:
            r = s.get(RANK_URL,
                      params={"eRank": "allex", "eOrd": eord, "eSort": esort},
                      timeout=12)
            r.encoding = "utf-8"
            _, rows = _parse_rank_table(BeautifulSoup(r.text, "lxml"))
            new = 0
            for row in rows:
                code = _extract_code(row)
                if code and code not in all_rows_by_code:
                    all_rows_by_code[code] = row
                    new += 1
            total_requests += 1
            if new > 0:
                print(f"  [排序:{label:8s}] {len(rows)} 筆, 新增 {new} → 累計 {len(all_rows_by_code)}")
            time.sleep(0.35)
        except Exception as e:
            print(f"  [排序:{label}] 錯誤：{e}")

    print(f"\n  共 {total_requests} 次請求，最終抓到 {len(all_rows_by_code)} 檔不重複 ETF")

    # 整理成有序清單（補上「代碼」欄保證存在）
    result = []
    for code, row in sorted(all_rows_by_code.items()):
        row["代碼"] = code
        result.append(row)
    return result


# ---------------------------------------------------------------------------
# 讀本機 DB
# ---------------------------------------------------------------------------
def load_db_etfs() -> list:
    if not DB_PATH.exists():
        print(f"⚠ 找不到資料庫：{DB_PATH}")
        print("  請先執行：python etf_db/twse_loader.py")
        return []
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    rows = conn.execute("""
        SELECT
            p.etf_code    AS 代號,
            p.etf_name    AS ETF名稱,
            p.issuer      AS 發行商,
            p.category    AS 分類,
            p.etf_type    AS ETF類型,
            p.region      AS 投資地區,
            p.tracking_index AS 追蹤指數,
            p.dividend_type  AS 配息方式,
            p.expense_ratio  AS 費用率,
            p.aum_billion    AS 規模億,
            p.inception_date AS 掛牌日,
            p.currency       AS 計價幣別,
            s.close_price    AS 收盤價,
            s.ytd_return     AS 今年報酬率,
            s.snapshot_date  AS 行情日期
        FROM etf_profile p
        LEFT JOIN (
            SELECT etf_code, close_price, ytd_return, snapshot_date
            FROM etf_market_snapshot
            WHERE (etf_code, snapshot_date) IN (
                SELECT etf_code, MAX(snapshot_date)
                FROM etf_market_snapshot GROUP BY etf_code
            )
        ) s ON p.etf_code = s.etf_code
        ORDER BY p.etf_code
    """).fetchall()
    conn.close()
    result = [dict(r) for r in rows]
    return result


# ---------------------------------------------------------------------------
# 寫 xlsx
# ---------------------------------------------------------------------------
def _hdr(cell, bg="1F4E79"):
    cell.fill = PatternFill("solid", fgColor=bg)
    cell.font = Font(bold=True, color="FFFFFF")
    cell.alignment = Alignment(horizontal="center", vertical="center")


def _auto_width(ws):
    # 略過合併儲存格（MergedCell 沒有 column_letter），以真實儲存格算欄寬
    widths: dict = {}
    for row in ws.iter_rows():
        for cell in row:
            if isinstance(cell, MergedCell):
                continue
            n = len(str(cell.value or ""))
            if n > widths.get(cell.column_letter, 0):
                widths[cell.column_letter] = n
    for col_letter, max_len in widths.items():
        ws.column_dimensions[col_letter].width = min(max_len + 2, 42)


def write_xlsx(moneydj: list, db: list, path: Path):
    wb = openpyxl.Workbook()

    # ── Sheet 1：MoneyDJ 清單 ──────────────────────────────
    ws1 = wb.active
    ws1.title = "MoneyDJ清單"
    ws1.freeze_panes = "A2"
    if moneydj:
        cols = list(moneydj[0].keys())
        # 確保「代碼」在最前
        if "代碼" in cols:
            cols = ["代碼"] + [c for c in cols if c != "代碼"]
        ws1.append(cols)
        for cell in ws1[1]:
            _hdr(cell, "1F4E79")
        for row in moneydj:
            ws1.append([row.get(c, "") for c in cols])
    else:
        ws1.append(["無資料（MoneyDJ 抓取失敗）"])
    _auto_width(ws1)

    # ── Sheet 2：本機 DB 清單 ──────────────────────────────
    ws2 = wb.create_sheet("本機DB清單")
    ws2.freeze_panes = "A2"
    if db:
        cols2 = list(db[0].keys())
        ws2.append(cols2)
        for cell in ws2[1]:
            _hdr(cell, "1E4620")
        for row in db:
            ws2.append([row.get(c, "") for c in cols2])
    else:
        ws2.append(["無資料（請先執行 twse_loader.py）"])
    _auto_width(ws2)

    # ── Sheet 3：比對結果 ──────────────────────────────────
    ws3 = wb.create_sheet("比對結果")
    ws3.freeze_panes = "A3"

    mj_codes = {r.get("代碼", "") for r in moneydj if r.get("代碼")}
    db_codes = {r.get("代號", "") for r in db if r.get("代號")}
    mj_name  = {r.get("代碼", ""): r.get("ETF名稱", "") for r in moneydj}
    db_name  = {r.get("代號", ""): r.get("ETF名稱", "") for r in db}
    db_cat   = {r.get("代號", ""): r.get("分類", "")   for r in db}

    both        = sorted(mj_codes & db_codes)
    only_mj     = sorted(mj_codes - db_codes)
    only_db     = sorted(db_codes - mj_codes)

    # 摘要行
    ws3.merge_cells("A1:E1")
    ws3["A1"] = (
        f"MoneyDJ {len(mj_codes)} 檔 vs 本機DB {len(db_codes)} 檔　｜　"
        f"兩邊都有：{len(both)}　｜　"
        f"只在MoneyDJ（DB未收錄）：{len(only_mj)}　｜　"
        f"只在DB（MoneyDJ未列）：{len(only_db)}"
    )
    ws3["A1"].font = Font(bold=True, size=11)
    ws3["A1"].alignment = Alignment(horizontal="left", vertical="center")
    ws3.row_dimensions[1].height = 22

    # 欄位標題
    hdr_row = ["狀態", "代號", "MoneyDJ名稱", "DB名稱", "DB分類"]
    ws3.append(hdr_row)
    for cell in ws3[2]:
        _hdr(cell, "7B2C2C")

    # 填資料（兩邊都有 → 只在MJ → 只在DB）
    green  = PatternFill("solid", fgColor="E2EFDA")
    yellow = PatternFill("solid", fgColor="FFF2CC")
    blue   = PatternFill("solid", fgColor="DDEEFF")

    def add_row(status, code, mj_n, db_n, cat, fill):
        row_idx = ws3.max_row + 1
        ws3.append([status, code, mj_n, db_n, cat])
        for cell in ws3[row_idx]:
            cell.fill = fill

    for code in both:
        add_row("✅ 兩邊都有", code, mj_name.get(code, ""), db_name.get(code, ""), db_cat.get(code, ""), green)
    for code in only_mj:
        add_row("⚠ 只在MoneyDJ", code, mj_name.get(code, ""), "", "", yellow)
    for code in only_db:
        add_row("🔵 只在DB", code, "", db_name.get(code, ""), db_cat.get(code, ""), blue)

    _auto_width(ws3)

    # 儲存
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    print(f"\n✅ 已儲存：{path}")
    print(f"   MoneyDJ 清單：{len(mj_codes)} 檔")
    print(f"   本機DB 清單：{len(db_codes)} 檔")
    print(f"   ✅ 兩邊都有：{len(both)}")
    print(f"   ⚠ 只在MoneyDJ（DB漏收）：{len(only_mj)}")
    print(f"   🔵 只在DB（MoneyDJ無列）：{len(only_db)}")


# ---------------------------------------------------------------------------
# main
# ---------------------------------------------------------------------------
def main():
    print("=" * 55)
    print("  MoneyDJ vs 本機DB ETF 清單比對")
    print("=" * 55)

    print(f"\n[1/3] 抓取 MoneyDJ 台股ETF受益人數排行...")
    moneydj = fetch_moneydj_etfs()

    print(f"\n[2/3] 讀取本機 DB（{DB_PATH}）...")
    db = load_db_etfs()
    print(f"  → DB 共 {len(db)} 檔")

    print(f"\n[3/3] 建立 {OUT_PATH} ...")
    write_xlsx(moneydj, db, OUT_PATH)


if __name__ == "__main__":
    main()
